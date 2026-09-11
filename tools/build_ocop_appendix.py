from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "ocop" / "ocop_records.json"
MANIFEST_PATH = ROOT / "data" / "ocop" / "source_manifest.csv"
OUTPUT_PATH = ROOT / "docs" / "PHU_LUC_NGUON_DU_LIEU_OCOP_2025_2026.docx"

GREEN = "006045"
GREEN_DARK = "004F3B"
GREEN_LIGHT = "DDF4EA"
GOLD = "FFB900"
NAVY = RGBColor(15, 23, 43)
SLATE = RGBColor(98, 116, 142)
WHITE = RGBColor(255, 255, 255)
LIGHT_BORDER = "D9E2EC"
FONT_NAME = "Segoe UI"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LIGHT_BORDER, size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_keep_together(row) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False


def set_cell_width(cell, width: float) -> None:
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=9, bold=False, color=NAVY, italic=False) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_paragraph(document, text: str = "", *, size=10, bold=False, color=NAVY,
                  italic=False, align=None, before=0, after=5, line=1.12):
    paragraph = document.add_paragraph()
    if text:
        set_run_font(paragraph.add_run(text), size=size, bold=bold, color=color, italic=italic)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    return paragraph


def add_heading(document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    size = 16 if level == 1 else 12
    color = RGBColor(0, 96, 69) if level == 1 else RGBColor(15, 23, 43)
    set_run_font(paragraph.add_run(text), size=size, bold=True, color=color)
    if level == 1:
        p_pr = paragraph._p.get_or_add_pPr()
        bottom = OxmlElement("w:pBdr")
        border = OxmlElement("w:bottom")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:color"), GREEN)
        border.set(qn("w:space"), "4")
        bottom.append(border)
        p_pr.append(bottom)


def add_bullet(document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(text), size=9.5)


def add_table(document, headers: list[str], rows: list[list[object]], widths: list[float],
              font_size: float = 8.2):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(str(text)), size=8, bold=True, color=WHITE)
        set_cell_margins(cell)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_keep_together(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            set_cell_width(cell, widths[column_index])
            if row_index % 2 == 1:
                set_cell_shading(cell, "F7FAF9")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            set_run_font(paragraph.add_run("" if value is None else str(value)), size=font_size)
            set_cell_margins(cell)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=8, color=SLATE)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def format_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.add_run("OCOP LÂM ĐỒNG · PHỤ LỤC DỮ LIỆU"), size=8, bold=True, color=RGBColor(0, 96, 69))
    add_page_number(section.footer.paragraphs[0])
    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(10)
    normal.font.color.rgb = NAVY


def break_url(url: str) -> str:
    for token in ("/", "?", "&", "=", "-"):
        url = url.replace(token, token + "\u200b")
    return url


dataset = json.loads(DATA_PATH.read_text(encoding="utf-8"))
with MANIFEST_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
    sources = list(csv.DictReader(handle))

products = []
for batch in dataset["batches"]:
    for product in batch["products"]:
        products.append({**product, "batch": batch})

category_counts = Counter(product["category"] for product in products)
star_counts = Counter(product["batch"]["star_rating"] for product in products)
year_counts = Counter(product["batch"]["cert_year"] for product in products)
region_counts = Counter(product["batch"]["legacy_region"] for product in products)
level_counts = Counter(product["batch"]["verification_level"] for product in products)
subject_count = len({product["subject_name"] for product in products})

document = Document()
format_document(document)

# Trang bìa
cover = document.add_table(rows=1, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cover.autofit = False
cell = cover.cell(0, 0)
set_cell_shading(cell, GREEN_DARK)
set_cell_margins(cell, top=180, start=180, bottom=180, end=180)
cell.text = ""
p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("TRƯỜNG ĐẠI HỌC ĐÀ LẠT\nKHOA CÔNG NGHỆ THÔNG TIN"), size=11, bold=True, color=WHITE)

for _ in range(4):
    document.add_paragraph()
add_paragraph(document, "PHỤ LỤC NGUỒN DỮ LIỆU", size=24, bold=True, color=RGBColor(0, 79, 59), align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
add_paragraph(document, "SẢN PHẨM OCOP LÂM ĐỒNG 2025–2026", size=19, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
add_paragraph(document, "Phạm vi tỉnh Lâm Đồng sau sắp xếp, gồm các khu vực Lâm Đồng, Bình Thuận và Đắk Nông trước đây", size=11, color=SLATE, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

summary = document.add_table(rows=3, cols=2)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
set_table_borders(summary, color=GREEN_LIGHT)
for row, (label, value) in zip(summary.rows, [
    ("Ngày chốt dữ liệu", dataset["metadata"]["cutoff_date"]),
    ("Quy mô", f"{len(products)} sản phẩm · {subject_count} chủ thể · {len(sources)} nguồn"),
    ("Mục đích", "Dữ liệu mẫu cho phát triển, kiểm thử và báo cáo đồ án"),
]):
    set_cell_width(row.cells[0], 1.8)
    set_cell_width(row.cells[1], 4.8)
    set_cell_shading(row.cells[0], GREEN_LIGHT)
    for index, text in enumerate((label, value)):
        row.cells[index].text = ""
        set_cell_margins(row.cells[index], top=110, start=120, bottom=110, end=120)
        set_run_font(row.cells[index].paragraphs[0].add_run(text), size=9.5, bold=index == 0, color=RGBColor(0, 79, 59) if index == 0 else NAVY)

add_paragraph(document, dataset["metadata"]["disclaimer"], size=9, bold=True, color=RGBColor(164, 38, 44), align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=0)
document.add_page_break()

# 1. Phạm vi và phương pháp
add_heading(document, "1. Phạm vi và mục đích", 1)
add_paragraph(document, "Phụ lục mô tả nguồn, phương pháp chuẩn hóa và danh sách sản phẩm được sử dụng trong cơ sở dữ liệu mẫu của Cổng thông tin OCOP và bản đồ du lịch nông nghiệp Lâm Đồng. Dữ liệu phục vụ học tập, phát triển phần mềm và kiểm thử; không thay thế hồ sơ công nhận do cơ quan nhà nước quản lý.")
add_bullet(document, "Phạm vi địa lý: tỉnh Lâm Đồng sau sắp xếp cấp tỉnh năm 2025, bao gồm ba khu vực Lâm Đồng, Bình Thuận và Đắk Nông trước đây.")
add_bullet(document, "Phạm vi nguồn: tài liệu công bố từ 01/01/2025 đến ngày 10/09/2026; năm chứng nhận được giữ nguyên theo nội dung nguồn.")
add_bullet(document, "Quy mô: 60 sản phẩm có nguồn cấp A hoặc B1; các nguồn B2 chỉ hỗ trợ đối chiếu tên, không tự xác nhận hạng sao.")
add_bullet(document, "Địa chỉ: giữ nguyên địa chỉ gốc; địa chỉ hiện hành chỉ được ghi khi đã ánh xạ theo Nghị quyết 1671/NQ-UBTVQH15.")
add_bullet(document, "Hình ảnh: chưa sử dụng ảnh báo chí hoặc ảnh chủ thể khi chưa có quyền sử dụng rõ ràng; giao diện dùng placeholder.")

add_heading(document, "2. Cấp độ tin cậy của nguồn", 1)
add_table(document,
          ["Cấp", "Mô tả", "Vai trò trong dữ liệu", "Điều kiện công khai"],
          [
              ["A", "Quyết định chính thức có chữ ký và phụ lục sản phẩm.", "Xác nhận tên, chủ thể, hạng sao, ngày quyết định.", "Được phép."],
              ["B1", "Báo/cổng cơ quan nhà nước xác nhận đã công nhận hoặc trao giấy.", "Tạo bản ghi đã xác nhận theo đúng nội dung bài.", "Có thể; ghi rõ trường còn thiếu."],
              ["B2", "Nguồn mới ở bước đề xuất, đánh giá hoặc chấm điểm.", "Đối chiếu định danh; không tự xác nhận kết quả.", "Không; trạng thái pending nếu dùng độc lập."],
              ["C", "Website của chủ thể hoặc nhà sản xuất.", "Bổ sung mô tả, giá, quy cách hoặc ảnh.", "Không dùng để tự xác nhận sao."],
          ],
          [0.55, 2.35, 2.45, 1.75], font_size=8.5)

add_heading(document, "3. Tổng quan bộ dữ liệu", 1)
add_table(document,
          ["Chỉ tiêu", "Số lượng", "Ghi chú"],
          [
              ["Sản phẩm", len(products), "Mục tiêu 60, vượt ngưỡng tối thiểu 50."],
              ["Chủ thể", subject_count, "Tài khoản kỹ thuật .invalid, bị khóa khi seed."],
              ["Nguồn", len(sources), "10 PDF lưu cục bộ; trang web được lưu URL."],
              ["Cấp A / B1", f"{level_counts['A']} / {level_counts['B1']}", "Nguồn chính của 60 bản ghi."],
              ["3 sao / 4 sao", f"{star_counts[3]} / {star_counts[4]}", "Không có sản phẩm 5 sao trong tập mẫu."],
              ["Năm chứng nhận", ", ".join(f"{year}: {year_counts[year]}" for year in sorted(year_counts)), "Ba sản phẩm năm 2024 được nguồn công bố trong năm 2025."],
          ],
          [2.0, 1.35, 3.75], font_size=8.7)

add_heading(document, "3.1. Phân bố theo khu vực nguồn", 2)
add_table(document, ["Khu vực", "Số sản phẩm"], [[name, count] for name, count in region_counts.items()], [5.6, 1.5], font_size=9)
add_heading(document, "3.2. Phân bố theo nhóm sản phẩm", 2)
add_table(document, ["Nhóm", "Số sản phẩm"], [[name, count] for name, count in category_counts.items()], [5.6, 1.5], font_size=9)
document.add_page_break()

# 4. Nguồn pháp lý và nguồn sản phẩm
add_heading(document, "4. Danh mục nguồn pháp lý bắt buộc", 1)
legal_sources = [source for source in sources if source["source_type"] == "legal_document"]
add_table(document,
          ["Mã nguồn", "Văn bản", "Cơ quan", "Ngày", "Tệp cục bộ"],
          [[source["source_id"], f"{source['document_number']} – {source['title']}", source["issuing_body"], source["published_at"], source["local_file"]] for source in legal_sources],
          [1.2, 2.55, 1.25, 0.75, 1.55], font_size=7.7)

add_heading(document, "5. Danh mục nguồn sản phẩm và nguồn đối chiếu", 1)
product_sources = [source for source in sources if source["source_type"] != "legal_document"]
add_table(document,
          ["Mã nguồn", "Tiêu đề", "Ngày", "Cấp", "Loại"],
          [[source["source_id"], source["title"], source["published_at"], source["verification_level"], source["source_type"]] for source in product_sources],
          [1.45, 3.35, 0.85, 0.55, 0.9], font_size=7.9)
document.add_page_break()

# 6. Danh sách sản phẩm theo nguồn
add_heading(document, "6. Danh sách 60 sản phẩm đã chuẩn hóa", 1)
add_paragraph(document, "Các tên dưới đây được giữ theo văn bản hoặc bài công bố. Trường chưa có quyết định, ngày cấp hoặc ngày hết hạn vẫn để trống trong Excel và PostgreSQL; không lấy ngày đăng bài thay cho ngày quyết định.", size=9.5, color=SLATE)
running_number = 1
for batch_index, batch in enumerate(dataset["batches"]):
    if batch_index > 0:
        document.add_page_break()
    source = next(source for source in sources if source["source_id"] == batch["source_id"])
    add_heading(document, f"6.{batch_index + 1}. {source['title']}", 2)
    add_paragraph(document,
                  f"Nguồn {batch['verification_level']} · {batch['star_rating']} sao · "
                  f"{batch['decision_number'] or 'chưa có số quyết định'} · "
                  f"{batch['cert_issued_at'] or 'chưa xác định ngày cấp'}",
                  size=8.8, bold=True, color=RGBColor(0, 96, 69), after=4)
    if batch.get("batch_notes"):
        add_paragraph(document, batch["batch_notes"], size=8.2, italic=True, color=SLATE, after=5)
    rows = []
    for product in batch["products"]:
        rows.append([
            running_number,
            product["product_name"],
            product["subject_name"],
            batch["star_rating"],
            product["current_address"],
        ])
        running_number += 1
    add_table(document, ["STT", "Sản phẩm", "Chủ thể", "Sao", "Địa chỉ hiện hành"], rows, [0.42, 2.0, 2.25, 0.45, 2.0], font_size=7.8)

document.add_page_break()

# 7. Ánh xạ hành chính
add_heading(document, "7. Ánh xạ địa chỉ trước và sau sắp xếp", 1)
add_paragraph(document, "Bảng ánh xạ chỉ bao gồm các xã/phường xuất hiện trong tập mẫu và đã được đối chiếu. Địa chỉ gốc vẫn được bảo toàn trong bảng product_sources để truy vết.", size=9.5)
mapping_rows = [[
    item["old_province"], item["old_district"], item["old_commune"],
    item["current_commune"], item["legal_source"],
] for item in dataset["admin_mappings"]]
add_table(document, ["Tỉnh cũ", "Huyện cũ", "Xã/phường cũ", "Xã/phường hiện tại", "Căn cứ"], mapping_rows, [1.05, 1.05, 1.7, 1.9, 1.4], font_size=8)

document.add_page_break()
add_heading(document, "8. Vấn đề còn mở và cách xử lý", 1)
add_table(document,
          ["Vấn đề", "Phạm vi", "Cách xử lý trong dữ liệu"],
          [
              ["Thiếu phụ lục quyết định", "Đức Trọng, Tánh Linh và nhóm 21 sản phẩm Bình Thuận cũ.", "Giữ nguyên B1, ghi chú và tiếp tục tìm quyết định; không tự tạo số."],
              ["Thiếu ngày cấp/hết hạn", "Các nhóm chỉ có bài trao giấy.", "Để NULL; không lấy ngày đăng bài làm ngày cấp và không tự cộng ba năm."],
              ["Thiếu cấp xã", "Rong nho và Rong nho tách nước.", "Giữ địa chỉ khái quát tỉnh Lâm Đồng; đánh dấu issue cần xác minh."],
              ["Năm chứng nhận 2024", "Ba sản phẩm Tánh Linh.", "Giữ cert_year=2024 vì bài năm 2025 mô tả đợt công nhận cuối năm 2024."],
              ["Ảnh chưa rõ quyền", "Toàn bộ 60 sản phẩm.", "Dùng placeholder cục bộ; chỉ thay khi chủ thể cung cấp hoặc có giấy phép."],
              ["Chưa có đủ 21 tên Bình Thuận", "Bài B1 xác nhận 21 sản phẩm nhưng không liệt kê đầy đủ.", "Chỉ nhập 11 tên xác định được từ nguồn B2 và đối chiếu kết quả bằng B1."],
          ],
          [1.65, 2.25, 3.5], font_size=8.4)

add_heading(document, "9. Quy tắc nạp dữ liệu vào hệ thống", 1)
add_bullet(document, "Seed sử dụng khóa tự nhiên và ON CONFLICT nên chạy lại không tạo sản phẩm, nguồn hoặc liên kết trùng.")
add_bullet(document, "Chín sản phẩm theo Quyết định 3981/QĐ-UBND giữ nguyên slug đã dùng trong seed_public_reference.sql.")
add_bullet(document, "Tài khoản kỹ thuật của chủ thể dùng email .invalid, is_active=false và không thể đăng nhập.")
add_bullet(document, "Giá không được nguồn công bố được lưu price=0 và unit='liên hệ chủ thể'.")
add_bullet(document, "API công khai chỉ lấy sản phẩm status='approved'; nguồn B2 độc lập không được tự động duyệt.")
add_bullet(document, "Số điện thoại cá nhân, giới tính và email người đại diện không được xuất hiện trong API công khai.")

document.add_page_break()
add_heading(document, "10. Danh mục URL nguồn", 1)
for index, source in enumerate(sources, start=1):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    set_run_font(p.add_run(f"{index}. {source['source_id']} · {source['title']}\n"), size=8.2, bold=True, color=RGBColor(0, 96, 69))
    set_run_font(p.add_run(break_url(source["source_url"])), size=7.4, color=SLATE)
    if source["local_file"]:
        set_run_font(p.add_run(f"\nTệp cục bộ: {source['local_file']} · SHA-256: {source['sha256']}"), size=7.2, color=SLATE)

add_paragraph(document, "Kết thúc phụ lục", size=9, bold=True, color=RGBColor(0, 96, 69), align=WD_ALIGN_PARAGRAPH.CENTER, before=14, after=0)

OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT_PATH)
print({"output": str(OUTPUT_PATH), "products": len(products), "sources": len(sources), "subjects": subject_count})
